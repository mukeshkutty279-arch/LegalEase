from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


def format_docx(document_text: str) -> bytes:
    """
    Convert generated legal document text into a DOCX file.
    """

    doc = Document()

    paragraphs = document_text.split("\n")

    for line in paragraphs:

        line = line.strip()

        if not line:
            doc.add_paragraph()
            continue

        paragraph = doc.add_paragraph()

        if (
            line.isupper()
            and len(line) < 100
        ):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            run = paragraph.add_run(line)
            run.bold = True
        else:
            paragraph.add_run(line)

    output = BytesIO()
    doc.save(output)

    return output.getvalue()
